import io
import os
import re
import shutil
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageEnhance, ImageFilter
import pdfplumber
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, portrait
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

MONTH_NAMES_ID = [
    "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember"
]

MONTH_MAP = {
    "JAN": "Januari", "FEB": "Februari", "MAR": "Maret", "APR": "April",
    "MAY": "Mei", "MEI": "Mei", "JUN": "Juni", "JUL": "Juli",
    "AUG": "Agustus", "AGU": "Agustus", "SEP": "September", "OCT": "Oktober",
    "OKT": "Oktober", "NOV": "November", "DEC": "Desember", "DES": "Desember"
}

def get_tesseract_cmd():
    """Mencari path binary tesseract pada sistem."""
    tess_path = shutil.which("tesseract")
    if not tess_path:
        for p in ["/opt/homebrew/bin/tesseract", "/usr/local/bin/tesseract", "/usr/bin/tesseract"]:
            if os.path.exists(p):
                return p
    return tess_path

def preprocess_image(img):
    """Meningkatkan kontras dan ketajaman teks angka untuk OCR."""
    gray = img.convert("L")
    enhancer = ImageEnhance.Contrast(gray)
    contrasted = enhancer.enhance(2.0)
    return contrasted.filter(ImageFilter.SHARPEN)

def detect_bca_month(full_text, filename=""):
    """Mendeteksi nama bulan transaksi Bank BCA dari teks atau nama berkas."""
    txt = (full_text + " " + filename).upper()
    
    # Format BCA: PERIODE : NOVEMBER 2025
    pm = re.search(r"PERIODE\s*:\s*([A-Za-z]+)\s+\d{4}", full_text, re.IGNORECASE)
    if pm:
        raw_m = pm.group(1).capitalize()
        if raw_m in MONTH_NAMES_ID:
            return raw_m
        for k, v in MONTH_MAP.items():
            if k in raw_m.upper():
                return v

    for m in MONTH_NAMES_ID:
        if m.upper() in txt:
            return m

    for code, name in MONTH_MAP.items():
        if re.search(rf"\b{code}\b", txt):
            return name

    return "Bulan"

def convert_pdf_to_images(pdf_input, dpi=300):
    """
    Konversi PDF ke list PIL Image (300 DPI).
    Mendukung pypdfium2 & pdf2image.
    """
    try:
        import pypdfium2
        if hasattr(pdf_input, "getvalue"):
            raw = pdf_input.getvalue()
            doc = pypdfium2.PdfDocument(raw)
        elif isinstance(pdf_input, bytes):
            doc = pypdfium2.PdfDocument(pdf_input)
        elif isinstance(pdf_input, (str, os.PathLike)) and os.path.exists(pdf_input):
            doc = pypdfium2.PdfDocument(str(pdf_input))
        elif hasattr(pdf_input, "read"):
            data = pdf_input.read()
            if hasattr(pdf_input, "seek"):
                pdf_input.seek(0)
            doc = pypdfium2.PdfDocument(data)
        else:
            doc = pypdfium2.PdfDocument(pdf_input)

        scale = dpi / 72.0
        images = [page.render(scale=scale).to_pil() for page in doc]
        if images:
            return images
    except Exception:
        pass

    try:
        from pdf2image import convert_from_path, convert_from_bytes
        poppler_path = None
        for p in ["/opt/homebrew/bin", "/usr/local/bin", "/usr/bin"]:
            if os.path.exists(os.path.join(p, "pdftoppm")):
                poppler_path = p
                break

        if isinstance(pdf_input, (str, os.PathLike)) and os.path.exists(pdf_input):
            return convert_from_path(str(pdf_input), dpi=dpi, poppler_path=poppler_path)
        elif hasattr(pdf_input, "getvalue"):
            return convert_from_bytes(pdf_input.getvalue(), dpi=dpi, poppler_path=poppler_path)
        elif isinstance(pdf_input, bytes):
            return convert_from_bytes(pdf_input.dpi, dpi=dpi, poppler_path=poppler_path)
        elif hasattr(pdf_input, "read"):
            data = pdf_input.read()
            if hasattr(pdf_input, "seek"):
                pdf_input.seek(0)
            return convert_from_bytes(data, dpi=dpi, poppler_path=poppler_path)
    except Exception:
        pass

    return []

def extract_ocr_text(images):
    """Mengekstrak teks dari daftar gambar menggunakan Tesseract OCR."""
    try:
        import pytesseract
        tess_cmd = get_tesseract_cmd()
        if tess_cmd:
            pytesseract.pytesseract.tesseract_cmd = tess_cmd

        try:
            installed_langs = pytesseract.get_languages()
            lang = "ind+eng" if "ind" in installed_langs else "eng"
        except Exception:
            lang = "eng"

        all_lines = []
        full_text = ""
        for img in images:
            processed = preprocess_image(img)
            try:
                txt = pytesseract.image_to_string(processed, lang=lang, config="--psm 6")
            except Exception:
                txt = pytesseract.image_to_string(processed, lang="eng")
            full_text += "\n" + txt
            for line in txt.split("\n"):
                l_str = line.strip()
                if l_str:
                    all_lines.append(l_str)
        return full_text, all_lines
    except Exception as e:
        print(f"[-] Error OCR: {e}")
        return "", []

def extract_pdf_digital_text(pdf_input):
    """Mengekstrak teks digital langsung dari PDF (akurasi 100% tanpa noise OCR)."""
    all_lines = []
    full_text = ""
    try:
        if hasattr(pdf_input, "getvalue"):
            stream = io.BytesIO(pdf_input.getvalue())
        elif isinstance(pdf_input, bytes):
            stream = io.BytesIO(pdf_input)
        elif isinstance(pdf_input, (str, os.PathLike)) and os.path.exists(pdf_input):
            stream = str(pdf_input)
        elif hasattr(pdf_input, "read"):
            data = pdf_input.read()
            if hasattr(pdf_input, "seek"):
                pdf_input.seek(0)
            stream = io.BytesIO(data)
        else:
            stream = pdf_input

        with pdfplumber.open(stream) as pdf:
            for p in pdf.pages:
                txt = p.extract_text() or ""
                full_text += "\n" + txt
                for line in txt.split("\n"):
                    l_str = line.strip()
                    if l_str:
                        all_lines.append(l_str)
    except Exception:
        pass
    return full_text, all_lines

def parse_bca(pdf_input, filename=""):
    """
    Engine Parser Rekening Koran Bank Central Asia (BCA):
    1. Ekstraksi teks digital atau fallback OCR 300 DPI.
    2. Deteksi bulan periode rekening giro/tahapan BCA.
    3. Ekstraksi ringkasan resmi BCA di footer:
       SALDO AWAL, MUTASI CR, MUTASI DB, SALDO AKHIR.
    4. Ekstraksi mutasi riil per baris (Tanggal DD/MM, Debet, Kredit, Saldo).
    """
    fname = filename or getattr(pdf_input, "name", "") or (str(pdf_input) if isinstance(pdf_input, str) else "rekening_bca.pdf")
    print(f"\n[+] Memproses berkas BCA: {os.path.basename(fname)} ...")

    all_lines = []
    full_text = ""

    # 1. Ekstraksi teks digital atau OCR fallback
    dig_text, dig_lines = extract_pdf_digital_text(pdf_input)
    if len(dig_text.strip()) > 100:
        full_text = dig_text
        all_lines = dig_lines
    else:
        print("  [*] Dokumen scan/gambar terdeteksi, memproses via Tesseract OCR 300 DPI...")
        images = convert_pdf_to_images(pdf_input, dpi=300)
        if images:
            full_text, all_lines = extract_ocr_text(images)

    bulan = detect_bca_month(full_text, fname)

    # 2. Ringkasan Resmi BCA (Footer)
    # SALDO AWAL : 40,358,423.11
    # MUTASI CR : 1,548,883,126.00 118
    # MUTASI DB : 1,562,800,646.00 312
    # SALDO AKHIR : 26,440,903.11
    awal_m = re.search(r"SALDO AWAL\s*:\s*([\d,\.]+)", full_text, re.IGNORECASE)
    cr_m = re.search(r"MUTASI CR\s*:\s*([\d,\.]+)\s*(\d+)?", full_text, re.IGNORECASE)
    db_m = re.search(r"MUTASI DB\s*:\s*([\d,\.]+)\s*(\d+)?", full_text, re.IGNORECASE)
    akhir_m = re.search(r"SALDO AKHIR\s*:\s*([\d,\.]+)", full_text, re.IGNORECASE)

    official_op = float(awal_m.group(1).replace(",", "")) if awal_m else None
    official_cr = float(cr_m.group(1).replace(",", "")) if cr_m else None
    official_f_cr = int(cr_m.group(2)) if (cr_m and cr_m.group(2)) else None

    official_db = float(db_m.group(1).replace(",", "")) if db_m else None
    official_f_db = int(db_m.group(2)) if (db_m and db_m.group(2)) else None

    # 3. Baris Transaksi Riil
    tx_records = []
    active_date = ""
    cur_saldo = official_op

    for line in all_lines:
        line_clean = line.strip()
        if any(x in line_clean.upper() for x in [
            "NO. REKENING", "HALAMAN", "PERIODE", "MATA UANG", "CATATAN",
            "BERSAMBUNG", "SALDO AWAL :", "SALDO AKHIR :", "MUTASI CR :",
            "MUTASI DB :", "TANGGAL KETERANGAN", "REKENING GIRO", "KCP MANGGA DUA"
        ]):
            continue

        # Cek tanggal transaksi format DD/MM di awal baris
        tgl_m = re.match(r"^(\d{2}/\d{2})\b", line_clean)
        if tgl_m:
            active_date = tgl_m.group(1)

        # Pola Debet: <nominal> DB [<saldo>]
        m_db = re.search(r"([\d,]+\.\d{2})\s+DB(?:\s+([\d,]+\.\d{2}))?", line_clean)
        if m_db:
            nom_db = float(m_db.group(1).replace(",", ""))
            s_val = float(m_db.group(2).replace(",", "")) if m_db.group(2) else None
            if s_val is not None:
                cur_saldo = s_val
            elif cur_saldo is not None:
                cur_saldo = round(cur_saldo - nom_db, 2)

            tx_records.append({
                "date": active_date,
                "debet": nom_db,
                "kredit": 0.0,
                "saldo": cur_saldo
            })
            continue

        # Pola Kredit
        if not re.search(r"TANGGAL\s*:", line_clean, re.IGNORECASE):
            amounts = re.findall(r"\b([\d,]+\.\d{2})\b", line_clean)
            if amounts:
                nums = [float(a.replace(",", "")) for a in amounts if float(a.replace(",", "")) < 500_000_000_000]
                if len(nums) >= 2 and "SALDO AWAL" not in line_clean.upper():
                    nom_cr, s_val = nums[-2], nums[-1]
                    cur_saldo = s_val
                    tx_records.append({
                        "date": active_date,
                        "debet": 0.0,
                        "kredit": nom_cr,
                        "saldo": s_val
                    })
                elif len(nums) == 1 and tgl_m and "SALDO AWAL" not in line_clean.upper():
                    nom_cr = nums[0]
                    if cur_saldo is not None:
                        cur_saldo = round(cur_saldo + nom_cr, 2)
                    tx_records.append({
                        "date": active_date,
                        "debet": 0.0,
                        "kredit": nom_cr,
                        "saldo": cur_saldo
                    })

    calc_db = round(sum(r["debet"] for r in tx_records), 2)
    calc_cr = round(sum(r["kredit"] for r in tx_records), 2)
    calc_f_db = sum(1 for r in tx_records if r["debet"] > 0)
    calc_f_cr = sum(1 for r in tx_records if r["kredit"] > 0)

    final_db = official_db if official_db is not None else calc_db
    final_cr = official_cr if official_cr is not None else calc_cr
    final_f_db = official_f_db if official_f_db is not None else calc_f_db
    final_f_cr = official_f_cr if official_f_cr is not None else calc_f_cr
    opening_bal = official_op if official_op is not None else (
        round(tx_records[0]["saldo"] + tx_records[0]["debet"] - tx_records[0]["kredit"], 2) if tx_records else 0.0
    )

    balances = [r["saldo"] for r in tx_records if r["saldo"] is not None and r["saldo"] >= 1000]
    saldo_max = max(balances) if balances else (opening_bal or 0.0)
    saldo_min = min(balances) if balances else (opening_bal or 0.0)
    saldo_avg = float(np.mean(balances)) if balances else (opening_bal or 0.0)

    print(f"  [✓] Sukses {bulan}: Debet={final_f_db}x (Rp {final_db:,.2f}) | Kredit={final_f_cr}x (Rp {final_cr:,.2f})")
    print(f"  [✓] Saldo {bulan}: Max=Rp {saldo_max:,.2f} | Avg=Rp {saldo_avg:,.2f} | Min=Rp {saldo_min:,.2f} | Total Baris={len(tx_records)}")

    return {
        "bank": "BCA",
        "bulan": bulan,
        "filename": fname,
        "opening_bal": opening_bal,
        "freq_db": final_f_db,
        "freq_cr": final_f_cr,
        "mutasi_db": final_db,
        "mutasi_cr": final_cr,
        "saldo_max": saldo_max,
        "saldo_avg": saldo_avg,
        "saldo_min": saldo_min,
        "tx_records": tx_records
    }

# Alias fungsi
parse_rekening_universal = parse_bca
parse_bca_clean = parse_bca

def sort_resume_chronological(resume_list):
    """Mengurutkan daftar resume rekening koran secara kronologis bulan Masehi."""
    def get_month_index(item):
        b = item.get("bulan", "")
        return MONTH_NAMES_ID.index(b) if b in MONTH_NAMES_ID else 99
    return sorted(resume_list, key=get_month_index)

# ==============================================================================
# GENERATOR FORM PDF
# ==============================================================================
def generate_form_pdf(output_target, header_info, resume_list, note_oh="", nama_so="", nama_oh=""):
    """
    Menghasilkan dokumen PDF Formulir Validasi Mutasi Rekening resmi BCA.
    Mendukung output ke path file (string) maupun in-memory BytesIO (return bytes).
    """
    sorted_resume = sort_resume_chronological(resume_list)

    is_buffer = isinstance(output_target, io.BytesIO) or output_target is None
    buffer = output_target if isinstance(output_target, io.BytesIO) else (io.BytesIO() if is_buffer else None)
    target = buffer if is_buffer else output_target

    doc = SimpleDocTemplate(
        target,
        pagesize=portrait(A4),
        rightMargin=18,
        leftMargin=18,
        topMargin=25,
        bottomMargin=25,
    )
    story = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=12,
        alignment=1,
        spaceAfter=15,
    )
    story.append(Paragraph("<b>FORM VALIDASI MUTASI REKENING</b>", title_style))

    header_rows = [
        ["Cabang", ":", header_info.get("cabang", "")],
        ["Nama Cust", ":", header_info.get("nama_cust", "")],
        ["", "", ""],
        ["Nomor Rekening", ":", header_info.get("no_rekening", "")],
        ["Nama Bank", ":", header_info.get("nama_bank", "BCA")],
        ["Nama Pemegang Rekening", ":", header_info.get("nama_pemegang_rek", "")],
    ]

    t_header = Table(header_rows, colWidths=[140, 15, 404])
    t_header.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BACKGROUND", (2, 0), (2, 1), colors.HexColor("#E2EFDA")),
        ("BACKGROUND", (2, 3), (2, 5), colors.HexColor("#E2EFDA")),
        ("BOX", (2, 0), (2, 1), 0.5, colors.black),
        ("BOX", (2, 3), (2, 5), 0.5, colors.black),
    ]))
    story.append(t_header)
    story.append(Spacer(1, 12))

    sec_title_style = ParagraphStyle(
        "SecTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9.5,
        spaceAfter=4,
    )
    story.append(Paragraph("<b>Resume Mutasi Rekening</b>", sec_title_style))

    table_data = [
        ["Bulan", "Frekuensi", "", "Mutasi", "", "Saldo", "", ""],
        ["", "Debet", "Kredit", "Debet", "Kredit", "Tertinggi", "Rata-Rata", "Terendah"],
    ]

    totals = {"f_db": 0, "f_cr": 0, "m_db": 0.0, "m_cr": 0.0, "s_max": 0.0, "s_avg": 0.0, "s_min": 0.0}
    n = len(sorted_resume)

    for item in sorted_resume:
        totals["f_db"] += item.get("freq_db", 0)
        totals["f_cr"] += item.get("freq_cr", 0)
        totals["m_db"] += item.get("mutasi_db", 0.0)
        totals["m_cr"] += item.get("mutasi_cr", 0.0)
        totals["s_max"] += item.get("saldo_max", 0.0)
        totals["s_avg"] += item.get("saldo_avg", 0.0)
        totals["s_min"] += item.get("saldo_min", 0.0)

        table_data.append([
            item.get("bulan", "-"),
            f"{item.get('freq_db', 0):,}",
            f"{item.get('freq_cr', 0):,}",
            f"{item.get('mutasi_db', 0.0):,.2f}",
            f"{item.get('mutasi_cr', 0.0):,.2f}",
            f"{item.get('saldo_max', 0.0):,.2f}",
            f"{item.get('saldo_avg', 0.0):,.2f}",
            f"{item.get('saldo_min', 0.0):,.2f}",
        ])

    table_data.append([
        "Rata-Rata",
        f"{int(round(totals['f_db']/n)):,}" if n else "0",
        f"{int(round(totals['f_cr']/n)):,}" if n else "0",
        f"{totals['m_db']/n:,.2f}" if n else "0.00",
        f"{totals['m_cr']/n:,.2f}" if n else "0.00",
        f"{totals['s_max']/n:,.2f}" if n else "0.00",
        f"{totals['s_avg']/n:,.2f}" if n else "0.00",
        f"{totals['s_min']/n:,.2f}" if n else "0.00",
    ])

    col_widths = [56, 38, 38, 86, 86, 85, 85, 85]
    t_resume = Table(table_data, colWidths=col_widths)
    t_resume.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#D9E1F2")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTNAME", (0, 0), (-1, 1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.2),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("SPAN", (0, 0), (0, 1)),
        ("SPAN", (1, 0), (2, 0)),
        ("SPAN", (3, 0), (4, 0)),
        ("SPAN", (5, 0), (7, 0)),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E2EFDA")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("ALIGN", (3, 2), (-1, -1), "RIGHT"),
        ("RIGHTPADDING", (3, 2), (-1, -1), 3),
        ("LEFTPADDING", (3, 2), (-1, -1), 3),
    ]))
    story.append(t_resume)
    story.append(Spacer(1, 15))

    story.append(Paragraph("<b>Note OH:</b>", sec_title_style))
    note_box = Table([[note_oh if note_oh else "-"]], colWidths=[559])
    note_box.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.5, colors.black),
        ("MINROWHEIGHT", (0, 0), (-1, -1), 35),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(note_box)
    story.append(Spacer(1, 30))

    so_display = f"SO: {nama_so}" if nama_so else "SO:"
    oh_display = f"Operation Head: {nama_oh}" if nama_oh else "Operation Head:"
    sign_data = [
        ["Mengajukan,", "Menyetujui,"],
        ["", ""],
        ["", ""],
        ["", ""],
        [so_display, oh_display],
    ]
    t_sign = Table(sign_data, colWidths=[275, 284])
    t_sign.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(t_sign)

    doc.build(story)
    print(f"[✓] Form PDF selesai dibuat.")

    if is_buffer:
        buffer.seek(0)
        return buffer.getvalue()
    return target

# ==============================================================================
# GENERATOR FORM EXCEL
# ==============================================================================
def generate_form_excel(output_target, header_info, resume_list, note_oh="", nama_so="", nama_oh=""):
    """
    Menghasilkan dokumen Excel Spreadsheet Formulir Validasi Mutasi Rekening resmi BCA.
    Lengkap dengan rincian transaksi per bulan di kolom sebelah kanan.
    """
    sorted_resume = sort_resume_chronological(resume_list)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Validasi Mutasi"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Arial", size=11, bold=True)
    font_sec = Font(name="Arial", size=10, bold=True)
    font_bold = Font(name="Arial", size=9, bold=True)
    font_norm = Font(name="Arial", size=9)

    fill_green = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_blue = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    fill_grey = PatternFill(start_color="808080", end_color="808080", fill_type="solid")
    fill_soft_blue = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")

    thin = Side(border_style="thin", color="000000")
    box_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # 1. Header Identitas Nasabah
    ws.merge_cells("C1:G1")
    ws["C1"] = "FORM VALIDASI MUTASI REKENING"
    ws["C1"].font = font_title
    ws["C1"].alignment = Alignment(horizontal="center", vertical="center")

    headers = [
        (3, "Cabang", header_info.get("cabang", "")),
        (4, "Nama Cust", header_info.get("nama_cust", "")),
        (6, "Nomor Rekening", header_info.get("no_rekening", "")),
        (7, "Nama Bank", header_info.get("nama_bank", "BCA")),
        (8, "Nama Pemegang Rekening", header_info.get("nama_pemegang_rek", "")),
    ]

    for r, label, val in headers:
        ws.cell(row=r, column=1, value=label).font = font_norm
        ws.cell(row=r, column=3, value=":").font = font_norm
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=7)
        cell = ws.cell(row=r, column=4, value=val)
        cell.font = font_norm
        cell.fill = fill_green
        for c in range(4, 8):
            ws.cell(row=r, column=c).border = box_border

    # 2. Tabel Resume Mutasi Kiri
    ws["A10"] = "Resume Mutasi Rekening"
    ws["A10"].font = font_sec

    ws.merge_cells("A11:A12")
    ws["A11"] = "Bulan"
    ws.merge_cells("B11:C11")
    ws["B11"] = "Frekuensi"
    ws["B12"] = "Debet"
    ws["C12"] = "Kredit"
    ws.merge_cells("D11:E11")
    ws["D11"] = "Mutasi"
    ws["D12"] = "Debet"
    ws["E12"] = "Kredit"
    ws.merge_cells("F11:H11")
    ws["F11"] = "Saldo"
    ws["F12"] = "Tertinggi"
    ws["G12"] = "Rata-Rata"
    ws["H12"] = "Terendah"

    for r in range(11, 13):
        for c in range(1, 9):
            cell = ws.cell(row=r, column=c)
            cell.fill = fill_blue
            cell.font = font_bold
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = box_border

    start_r = 13
    n = len(sorted_resume)
    sum_vals = {"f_db": 0, "f_cr": 0, "m_db": 0.0, "m_cr": 0.0, "s_max": 0.0, "s_avg": 0.0, "s_min": 0.0}

    for i, item in enumerate(sorted_resume):
        curr_r = start_r + i
        b_name = item.get("bulan", "-")

        c1 = ws.cell(row=curr_r, column=1, value=b_name)
        c1.alignment = Alignment(horizontal="center")
        c1.font = font_norm
        c1.border = box_border

        f_db = item.get("freq_db", 0)
        f_cr = item.get("freq_cr", 0)
        m_db = item.get("mutasi_db", 0.0)
        m_cr = item.get("mutasi_cr", 0.0)
        s_max = item.get("saldo_max", 0.0)
        s_avg = item.get("saldo_avg", 0.0)
        s_min = item.get("saldo_min", 0.0)

        sum_vals["f_db"] += f_db
        sum_vals["f_cr"] += f_cr
        sum_vals["m_db"] += m_db
        sum_vals["m_cr"] += m_cr
        sum_vals["s_max"] += s_max
        sum_vals["s_avg"] += s_avg
        sum_vals["s_min"] += s_min

        ws.cell(row=curr_r, column=2, value=f_db).number_format = "#,##0"
        ws.cell(row=curr_r, column=3, value=f_cr).number_format = "#,##0"
        ws.cell(row=curr_r, column=4, value=m_db).number_format = "#,##0.00"
        ws.cell(row=curr_r, column=5, value=m_cr).number_format = "#,##0.00"
        ws.cell(row=curr_r, column=6, value=s_max).number_format = "#,##0.00"
        ws.cell(row=curr_r, column=7, value=s_avg).number_format = "#,##0.00"
        ws.cell(row=curr_r, column=8, value=s_min).number_format = "#,##0.00"

        for c_idx in range(2, 9):
            cell = ws.cell(row=curr_r, column=c_idx)
            cell.font = font_norm
            cell.border = box_border
            cell.alignment = Alignment(horizontal="right")

    end_data_r = start_r + len(sorted_resume) - 1
    avg_r = end_data_r + 1

    ws.cell(row=avg_r, column=1, value="Rata-Rata")
    ws.cell(row=avg_r, column=2, value=int(round(sum_vals["f_db"] / n)) if n else 0).number_format = "#,##0"
    ws.cell(row=avg_r, column=3, value=int(round(sum_vals["f_cr"] / n)) if n else 0).number_format = "#,##0"
    ws.cell(row=avg_r, column=4, value=(sum_vals["m_db"] / n) if n else 0.0).number_format = "#,##0.00"
    ws.cell(row=avg_r, column=5, value=(sum_vals["m_cr"] / n) if n else 0.0).number_format = "#,##0.00"
    ws.cell(row=avg_r, column=6, value=(sum_vals["s_max"] / n) if n else 0.0).number_format = "#,##0.00"
    ws.cell(row=avg_r, column=7, value=(sum_vals["s_avg"] / n) if n else 0.0).number_format = "#,##0.00"
    ws.cell(row=avg_r, column=8, value=(sum_vals["s_min"] / n) if n else 0.0).number_format = "#,##0.00"

    for c in range(1, 9):
        cell = ws.cell(row=avg_r, column=c)
        cell.fill = fill_green
        cell.font = font_bold
        cell.border = box_border
        cell.alignment = Alignment(horizontal="center" if c == 1 else "right")

    note_label_r = avg_r + 2
    ws.cell(row=note_label_r, column=1, value="Note OH:").font = font_sec
    ws.merge_cells(start_row=note_label_r + 1, start_column=1, end_row=note_label_r + 2, end_column=8)
    note_c = ws.cell(row=note_label_r + 1, column=1, value=note_oh if note_oh else "-")
    note_c.font = font_norm
    note_c.alignment = Alignment(vertical="top")
    for r in range(note_label_r + 1, note_label_r + 3):
        for c in range(1, 9):
            ws.cell(row=r, column=c).border = box_border

    sign_r = note_label_r + 4
    ws.cell(row=sign_r, column=1, value="Mengajukan,").font = font_norm
    ws.cell(row=sign_r, column=6, value="Menyetujui,").font = font_norm
    ws.cell(row=sign_r + 4, column=1, value=f"SO: {nama_so}" if nama_so else "SO:").font = font_norm
    ws.cell(row=sign_r + 4, column=6, value=f"Operation Head: {nama_oh}" if nama_oh else "Operation Head:").font = font_norm

    # 3. Rincian Mutasi Kanan (Kolom J ke Kanan)
    start_col = 10
    for idx, item in enumerate(sorted_resume):
        b_name = item.get("bulan", f"Bulan {idx+1}")
        col_b = start_col + (idx * 5)

        ws.cell(row=1, column=col_b, value="Bulan")
        ws.merge_cells(start_row=1, start_column=col_b + 1, end_row=1, end_column=col_b + 2)
        ws.cell(row=1, column=col_b + 1, value=f"Mutasi {b_name}")
        ws.cell(row=1, column=col_b + 3, value="Saldo")

        ws.cell(row=2, column=col_b + 1, value="Debet")
        ws.cell(row=2, column=col_b + 2, value="Kredit")

        for r in range(1, 3):
            for c in range(col_b, col_b + 4):
                cell = ws.cell(row=r, column=c)
                cell.font = font_bold
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = box_border

        # Baris 3 Kuning: Total Mutasi Debet, Kredit, dan Saldo Awal Resmi
        ws.cell(row=3, column=col_b, value=b_name).fill = fill_yellow
        ws.cell(row=3, column=col_b).font = font_bold
        ws.cell(row=3, column=col_b).border = box_border

        ws.cell(row=3, column=col_b + 1, value=item.get("mutasi_db", 0.0))
        ws.cell(row=3, column=col_b + 2, value=item.get("mutasi_cr", 0.0))
        ws.cell(row=3, column=col_b + 3, value=item.get("opening_bal", 0.0) or 0.0)

        for c in range(col_b, col_b + 4):
            cell = ws.cell(row=3, column=c)
            cell.fill = fill_yellow
            cell.font = font_bold
            cell.border = box_border
            if c > col_b:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0.00"

        records = item.get("tx_records", [])
        max_tx_row = max(len(records) + 3, 40)

        # Baris 4 dst: Transaksi Mutasi Riil
        for r_idx in range(4, max_tx_row + 1):
            tx_idx = r_idx - 4
            rec = records[tx_idx] if tx_idx < len(records) else None

            c_bln = ws.cell(row=r_idx, column=col_b, value=rec["date"] if rec else "")
            c_bln.fill = fill_grey
            c_bln.font = Font(name="Arial", size=8, color="FFFFFF")
            c_bln.border = box_border
            c_bln.alignment = Alignment(horizontal="center")

            c_deb = ws.cell(row=r_idx, column=col_b + 1, value=rec["debet"] if (rec and rec["debet"] > 0) else None)
            c_deb.fill = fill_yellow
            c_deb.font = font_norm
            c_deb.border = box_border
            c_deb.number_format = "#,##0.00"

            c_krd = ws.cell(row=r_idx, column=col_b + 2, value=rec["kredit"] if (rec and rec["kredit"] > 0) else None)
            c_krd.fill = fill_yellow
            c_krd.font = font_norm
            c_krd.border = box_border
            c_krd.number_format = "#,##0.00"

            c_sal = ws.cell(row=r_idx, column=col_b + 3, value=rec["saldo"] if rec else None)
            c_sal.fill = fill_soft_blue
            c_sal.font = font_norm
            c_sal.border = box_border
            c_sal.number_format = "#,##0.00"
            c_sal.alignment = Alignment(horizontal="right")

        col_sep = col_b + 4
        for r_sep in range(1, max_tx_row + 1):
            ws.cell(row=r_sep, column=col_sep).fill = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
        ws.column_dimensions[get_column_letter(col_sep)].width = 2.5

        ws.column_dimensions[get_column_letter(col_b)].width = 12
        ws.column_dimensions[get_column_letter(col_b + 1)].width = 16
        ws.column_dimensions[get_column_letter(col_b + 2)].width = 16
        ws.column_dimensions[get_column_letter(col_b + 3)].width = 18

    for col in ["A", "B", "C", "D", "E", "F", "G", "H"]:
        ws.column_dimensions[col].width = 15
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 10

    is_buffer = isinstance(output_target, io.BytesIO) or output_target is None
    buffer = output_target if isinstance(output_target, io.BytesIO) else (io.BytesIO() if is_buffer else None)

    if is_buffer:
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
    else:
        wb.save(output_target)
        print(f"[✓] File Excel siap diunduh: {output_target}")
        return output_target
